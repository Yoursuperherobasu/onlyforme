import io
import mimetypes
import traceback
from pathlib import Path
from typing import Optional

from loguru import logger

from agentcore.custom.custom_node.node import Node
from agentcore.io import BoolInput, DropdownInput, HandleInput, IntInput, Output
from agentcore.schema.data import Data

# ═══════════════════════════════════════════════════════════════
#  HARDCODED API KEY — REPLACE WITH YOUR ACTUAL KEY
# ═══════════════════════════════════════════════════════════════
GEMINI_API_KEY = "AIzaSyC3UhBn_HLOEkvtbo1D8jhS58enFkaDjDo"
# ═══════════════════════════════════════════════════════════════


class GeminiOCRExtractorNode(Node):
    display_name: str = "Multimodal Document Loader"
    description: str = (
        "Extract text from files using Gemini Multimodal Vision. "
        "Supports PDF (native + scanned), images, DOCX, PPTX, XLSX, CSV, TXT."
    )
    name = "GeminiOCRExtractor"
    icon = "FileText"

    inputs = [
        DropdownInput(
            name="gemini_model",
            display_name="Gemini Model",
            options=["gemini-2.0-flash", "gemini-1.5-flash", "gemini-1.5-pro"],
            value="gemini-2.0-flash",
        ),
        HandleInput(
            name="file_paths",
            display_name="File Paths",
            input_types=["Data", "Message"],
            info="File paths from Knowledge Base component.",
            is_list=True,
        ),
        IntInput(name="min_native_text_length", display_name="Min Native Text Length", value=50, advanced=True),
        IntInput(name="ocr_dpi", display_name="OCR DPI", value=300, advanced=True),
        BoolInput(name="extract_tables", display_name="Extract Tables", value=True, advanced=True),
    ]

    outputs = [
        Output(
            display_name="Extracted Documents",
            name="documents",
            method="extract_documents",
            output_types=["Data"],
            is_list=True,
        ),
    ]

    SUPPORTED_EXTENSIONS = {
        ".pdf", ".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp",
        ".docx", ".pptx", ".xlsx", ".csv", ".txt",
    }
    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"}

    # ══════════════════════════════════════════════════════════
    #  MAIN ENTRY POINT
    # ══════════════════════════════════════════════════════════

    def extract_documents(self) -> list[Data]:
        import google.generativeai as genai

        genai.configure(api_key=GEMINI_API_KEY)
        self._model = genai.GenerativeModel(self.gemini_model)
        self._genai = genai

        # ── Debug: log raw input ──
        logger.info(f"[OCR] file_paths type: {type(self.file_paths)}")
        logger.info(f"[OCR] file_paths value: {self.file_paths}")

        if self.file_paths and isinstance(self.file_paths, list):
            for i, item in enumerate(self.file_paths):
                logger.info(f"[OCR] item[{i}] type={type(item).__name__}, has text={hasattr(item, 'text')}")
                if hasattr(item, "text"):
                    logger.info(f"[OCR] item[{i}].text = '{getattr(item, 'text', '')}'")

        # ── Resolve file paths ──
        paths = self._resolve_paths()
        logger.info(f"[OCR] Resolved {len(paths)} path(s): {paths}")

        if not paths:
            self.status = "No files found. Check Knowledge Base connection."
            return [Data(text="No files found to process.", data={"error": True})]

        # ── Extract text from each file ──
        all_docs: list[Data] = []
        errors: list[str] = []

        for path in paths:
            try:
                logger.info(f"[OCR] Extracting: {path} (exists={path.exists()})")
                docs = self._extract_file(path)
                logger.info(f"[OCR] Got {len(docs)} doc(s) from {path.name}")
                for d in docs:
                    logger.info(f"[OCR]   text length: {len(d.text) if d.text else 0}")
                all_docs.extend(docs)
            except Exception as e:
                err_msg = f"{path.name}: {type(e).__name__}: {e}"
                logger.error(f"[OCR] Error extracting {path}: {err_msg}")
                logger.error(traceback.format_exc())
                errors.append(err_msg)

        status = f"Extracted {len(all_docs)} document(s) from {len(paths)} file(s)"
        if errors:
            status += f" | Errors: {'; '.join(errors[:3])}"
        self.status = status
        logger.info(f"[OCR] Final status: {status}")

        # If extraction produced nothing, return error info
        if not all_docs:
            return [Data(
                text=f"Extraction returned 0 documents. Errors: {'; '.join(errors)}",
                data={"error": True, "paths": [str(p) for p in paths]},
            )]

        return all_docs

    # ══════════════════════════════════════════════════════════
    #  PATH RESOLUTION — type-agnostic, no isinstance on custom types
    #
    #  Knowledge Base load_files_path() returns Message with:
    #    .text = "C:/path/to/file.pptx"
    # ══════════════════════════════════════════════════════════

    def _resolve_paths(self) -> list[Path]:
        paths: list[Path] = []
        if not self.file_paths:
            logger.warning("[OCR] self.file_paths is empty/None")
            return paths

        items = self.file_paths if isinstance(self.file_paths, list) else [self.file_paths]
        logger.info(f"[OCR] Processing {len(items)} input item(s)")

        for item in items:
            if item is None:
                continue

            candidates: list[str] = []

            if isinstance(item, str):
                candidates.append(item)
            elif isinstance(item, dict):
                for key in ["file_path", "path", "file", "text", "source"]:
                    val = item.get(key)
                    if val and isinstance(val, str):
                        candidates.append(val)
            else:
                # ANY object — check .text FIRST (this is where KB puts the path)
                text_val = getattr(item, "text", None)
                if text_val and isinstance(text_val, str) and text_val.strip():
                    candidates.append(text_val.strip())
                    logger.info(f"[OCR] Found .text = '{text_val.strip()}'")

                # Check .data dict
                data_dict = getattr(item, "data", None)
                if data_dict and isinstance(data_dict, dict):
                    for key in ["file_path", "path", "file", "text", "source"]:
                        val = data_dict.get(key)
                        if val and isinstance(val, str) and val.strip():
                            candidates.append(val.strip())

                # Check .path attribute (only if non-None and non-empty)
                path_val = getattr(item, "path", None)
                if path_val is not None:
                    path_str = str(path_val).strip()
                    if path_str and path_str != "None":
                        candidates.append(path_str)

            # Validate all candidates
            for candidate in candidates:
                for line in candidate.split("\n"):
                    line = line.strip()
                    if not line or line == "None":
                        continue
                    try:
                        p = Path(line)
                        if p.suffix.lower() in self.SUPPORTED_EXTENSIONS:
                            if p.exists():
                                if p not in paths:
                                    paths.append(p)
                                    logger.info(f"[OCR] Valid path added: {p}")
                            else:
                                logger.warning(f"[OCR] Path does not exist: {p}")
                        else:
                            logger.debug(f"[OCR] Unsupported extension: {p.suffix}")
                    except (OSError, ValueError) as e:
                        logger.warning(f"[OCR] Invalid path '{line}': {e}")

        return paths

    # ══════════════════════════════════════════════════════════
    #  FILE DISPATCH
    # ══════════════════════════════════════════════════════════

    def _extract_file(self, path: Path) -> list[Data]:
        ext = path.suffix.lower()
        dispatch = {
            ".pdf": self._extract_pdf, ".docx": self._extract_docx,
            ".pptx": self._extract_pptx, ".xlsx": self._extract_xlsx,
            ".csv": self._extract_csv, ".txt": self._extract_txt,
        }
        if ext in dispatch:
            return dispatch[ext](path)
        elif ext in self.IMAGE_EXTENSIONS:
            return self._extract_image(path)
        return []

    def _extract_pdf(self, path: Path) -> list[Data]:
        from PyPDF2 import PdfReader
        docs = []
        reader = PdfReader(str(path))
        for page_num, page in enumerate(reader.pages, start=1):
            native_text = (page.extract_text() or "").strip()
            if len(native_text) >= self.min_native_text_length:
                docs.append(Data(text=native_text, data={
                    "source_file": path.name, "file_path": str(path),
                    "page_number": page_num, "total_pages": len(reader.pages),
                    "file_type": "pdf", "extraction_method": "native",
                }))
            else:
                ocr_text = self._ocr_pdf_page(path, page_num)
                if ocr_text:
                    docs.append(Data(text=ocr_text, data={
                        "source_file": path.name, "file_path": str(path),
                        "page_number": page_num, "total_pages": len(reader.pages),
                        "file_type": "pdf", "extraction_method": "gemini_vision",
                    }))
        return docs

    def _ocr_pdf_page(self, pdf_path: Path, page_number: int) -> Optional[str]:
        try:
            from pdf2image import convert_from_path
            images = convert_from_path(str(pdf_path), first_page=page_number, last_page=page_number, dpi=self.ocr_dpi)
            if not images:
                return None
            buf = io.BytesIO()
            images[0].save(buf, format="PNG")
            return self._gemini_vision_extract(buf.getvalue(), "image/png")
        except Exception:
            return None

    def _extract_image(self, path: Path) -> list[Data]:
        mime, _ = mimetypes.guess_type(str(path))
        text = self._gemini_vision_extract(path.read_bytes(), mime or "image/png")
        if text:
            return [Data(text=text, data={
                "source_file": path.name, "file_path": str(path),
                "page_number": 1, "file_type": "image", "extraction_method": "gemini_vision",
            })]
        return []

    def _extract_docx(self, path: Path) -> list[Data]:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        if self.extract_tables:
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        parts.append(row_text)
        full_text = "\n".join(parts)
        if full_text.strip():
            return [Data(text=full_text, data={
                "source_file": path.name, "file_path": str(path),
                "page_number": 1, "file_type": "docx", "extraction_method": "native",
            })]
        return []

    def _extract_pptx(self, path: Path) -> list[Data]:
        from pptx import Presentation
        logger.info(f"[OCR] Opening PPTX: {path}")
        prs = Presentation(str(path))
        logger.info(f"[OCR] PPTX has {len(prs.slides)} slides")
        docs = []
        for slide_num, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        if para.text.strip():
                            texts.append(para.text.strip())
                if self.extract_tables and shape.has_table:
                    for row in shape.table.rows:
                        rt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                        if rt:
                            texts.append(rt)
            if texts:
                slide_text = "\n".join(texts)
                logger.info(f"[OCR] Slide {slide_num}: {len(slide_text)} chars")
                docs.append(Data(text=slide_text, data={
                    "source_file": path.name, "file_path": str(path),
                    "page_number": slide_num, "total_pages": len(prs.slides),
                    "file_type": "pptx", "extraction_method": "native",
                }))
        return docs

    def _extract_xlsx(self, path: Path) -> list[Data]:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), data_only=True)
        docs = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                vals = [str(c) if c is not None else "" for c in row]
                rt = " | ".join(v for v in vals if v)
                if rt:
                    rows.append(rt)
            if rows:
                docs.append(Data(text="\n".join(rows), data={
                    "source_file": path.name, "file_path": str(path),
                    "sheet_name": sheet_name, "page_number": 1,
                    "file_type": "xlsx", "extraction_method": "native",
                }))
        return docs

    def _extract_csv(self, path: Path) -> list[Data]:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        if text:
            return [Data(text=text, data={
                "source_file": path.name, "file_path": str(path),
                "page_number": 1, "file_type": "csv", "extraction_method": "native",
            })]
        return []

    def _extract_txt(self, path: Path) -> list[Data]:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        if text:
            return [Data(text=text, data={
                "source_file": path.name, "file_path": str(path),
                "page_number": 1, "file_type": "txt", "extraction_method": "native",
            })]
        return []

    def _gemini_vision_extract(self, image_bytes: bytes, mime_type: str) -> Optional[str]:
        prompt = (
            "Extract ALL text from this image accurately. "
            "Preserve the original structure: headings, paragraphs, tables, lists. "
            "Format tables as markdown tables. "
            "Return ONLY the extracted text, no commentary or explanation."
        )
        try:
            response = self._model.generate_content(
                [prompt, {"mime_type": mime_type, "data": image_bytes}],
                generation_config=self._genai.types.GenerationConfig(temperature=0.0, max_output_tokens=4096),
            )
            text = response.text.strip()
            return text if text else None
        except Exception as e:
            logger.error(f"[OCR] Gemini vision error: {e}")
            return None