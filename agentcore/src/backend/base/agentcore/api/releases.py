from __future__ import annotations

import csv
import json
import re
from collections import defaultdict, deque
from datetime import date, datetime, timezone
from enum import Enum
from io import BytesIO, StringIO
from typing import Any
from uuid import UUID

from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError
from sqlalchemy.exc import ProgrammingError
from sqlmodel import select

from agentcore.api.utils import CurrentActiveUser, DbSession
from agentcore.services.auth.permissions import get_permissions_for_role
from agentcore.services.database.models.package.model import Package
from agentcore.services.database.models.product_release.model import ProductRelease
from agentcore.services.database.models.release_detail.model import ReleaseDetail
from agentcore.services.database.models.release_package_snapshot.model import ReleasePackageSnapshot

router = APIRouter(prefix="/releases", tags=["Release Management"])

ACTIVE_END_DATE = date(9999, 12, 31)


async def _require_release_permission(current_user: CurrentActiveUser, permission: str) -> None:
    if str(getattr(current_user, "role", "")).strip().lower() == "root":
        return
    user_permissions = await get_permissions_for_role(str(current_user.role))
    if permission not in user_permissions:
        raise HTTPException(status_code=403, detail="Missing required permissions.")


class BumpType(str, Enum):
    major = "major"
    minor = "minor"
    patch = "patch"


class ReleaseBumpRequest(BaseModel):
    bump_type: BumpType
    release_notes: str | None = None


class ReleaseDetailInput(BaseModel):
    section_no: int | None = None
    section_title: str | None = None
    module: str | None = None
    sub_module: str | None = None
    feature_capability: str
    description_details: str | None = None


def _parse_semver(version: str) -> tuple[int, int, int]:
    parts = version.split(".")
    if len(parts) != 3:
        raise ValueError(f"Invalid semantic version '{version}'. Expected format: X.Y.Z")
    return int(parts[0]), int(parts[1]), int(parts[2])


def _bump(major: int, minor: int, patch: int, bump_type: BumpType) -> tuple[int, int, int]:
    if bump_type == BumpType.major:
        return major + 1, 0, 0
    if bump_type == BumpType.minor:
        return major, minor + 1, 0
    return major, minor, patch + 1


def _release_to_payload(release: ProductRelease, package_count: int | None = None) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "id": str(release.id),
        "version": release.version,
        "major": release.major,
        "minor": release.minor,
        "patch": release.patch,
        "release_notes": release.release_notes or "",
        "start_date": release.start_date.isoformat(),
        "end_date": release.end_date.isoformat(),
        "created_by": str(release.created_by) if release.created_by else None,
        "created_at": release.created_at.isoformat(),
        "updated_at": release.updated_at.isoformat(),
        "is_active": release.end_date == ACTIVE_END_DATE,
    }
    if package_count is not None:
        payload["package_count"] = package_count
    return payload


def _detail_to_payload(detail: ReleaseDetail) -> dict[str, Any]:
    return {
        "id": str(detail.id),
        "release_id": str(detail.release_id),
        "section_no": detail.section_no,
        "section_title": detail.section_title,
        "module": detail.module,
        "sub_module": detail.sub_module,
        "feature_capability": detail.feature_capability,
        "description_details": detail.description_details,
        "sort_order": detail.sort_order,
        "created_at": detail.created_at.isoformat(),
    }


def _release_package_to_payload(snapshot: ReleasePackageSnapshot) -> dict[str, Any]:
    return {
        "id": str(snapshot.id),
        "release_id": str(snapshot.release_id),
        "service_name": snapshot.service_name,
        "name": snapshot.name,
        "version": snapshot.version,
        "version_spec": snapshot.version_spec,
        "package_type": snapshot.package_type,
        "required_by": snapshot.required_by or [],
        "source": snapshot.source or {},
        "captured_at": snapshot.captured_at.isoformat(),
    }


def _normalize_pkg(name: str) -> str:
    return re.sub(r"[-_.]+", "-", name.strip().lower())


def _extract_parent_name(raw_parent: str) -> str:
    parent = (raw_parent or "").strip()
    if not parent:
        return ""
    if ":" in parent:
        parent = parent.split(":", 1)[0].strip()
    return parent


def _build_release_parent_map(
    snapshots: list[ReleasePackageSnapshot],
) -> dict[str, set[str]]:
    parents_by_child: dict[str, set[str]] = {}
    for row in snapshots:
        if row.package_type != "transitive":
            continue
        child = _normalize_pkg(row.name)
        parents = {
            _normalize_pkg(parent_name)
            for parent_name in (_extract_parent_name(parent) for parent in (row.required_by or []))
            if parent_name
        }
        if parents:
            parents_by_child.setdefault(child, set()).update(parents)
    return parents_by_child


def _compute_reachable_transitives_from_packages(
    managed_rows: list[Package],
    transitive_rows: list[Package],
) -> set[str]:
    managed_names = {_normalize_pkg(row.name) for row in managed_rows}
    children_by_parent: dict[str, set[str]] = {}

    for row in transitive_rows:
        child = _normalize_pkg(row.name)
        parent_names: set[str] = set()
        for detail in row.required_by_details or []:
            parent = (detail.get("name") or "").strip()
            if parent:
                parent_names.add(_normalize_pkg(parent))
        for parent_name in (_extract_parent_name(parent) for parent in (row.required_by or [])):
            if parent_name:
                parent_names.add(_normalize_pkg(parent_name))
        for parent in parent_names:
            children_by_parent.setdefault(parent, set()).add(child)

    reachable: set[str] = set()
    queue = deque(managed_names)
    seen: set[str] = set(managed_names)
    while queue:
        parent = queue.popleft()
        for child in children_by_parent.get(parent, set()):
            if child in reachable:
                continue
            reachable.add(child)
            if child not in seen:
                seen.add(child)
                queue.append(child)
    return reachable


def _compute_reachable_transitives_from_snapshot(
    managed_rows: list[ReleasePackageSnapshot],
    transitive_rows: list[ReleasePackageSnapshot],
) -> set[str]:
    managed_names = {_normalize_pkg(row.name) for row in managed_rows}
    children_by_parent: dict[str, set[str]] = {}

    for row in transitive_rows:
        child = _normalize_pkg(row.name)
        parent_names: set[str] = set()
        for parent_name in (_extract_parent_name(parent) for parent in (row.required_by or [])):
            if parent_name:
                parent_names.add(_normalize_pkg(parent_name))
        for parent in parent_names:
            children_by_parent.setdefault(parent, set()).add(child)

    reachable: set[str] = set()
    queue = deque(managed_names)
    seen: set[str] = set(managed_names)
    while queue:
        parent = queue.popleft()
        for child in children_by_parent.get(parent, set()):
            if child in reachable:
                continue
            reachable.add(child)
            if child not in seen:
                seen.add(child)
                queue.append(child)
    return reachable


def _compute_managed_root_data_for_snapshot(
    package_name: str,
    parents_by_child: dict[str, set[str]],
    managed_names: set[str],
    display_names: dict[str, str],
) -> tuple[list[str], list[str]]:
    child = _normalize_pkg(package_name)
    roots: set[str] = set()
    paths: list[str] = []
    max_paths = 25

    stack: list[tuple[str, list[str]]] = [(child, [child])]
    seen_states: set[tuple[str, tuple[str, ...]]] = set()

    while stack and len(paths) < max_paths:
        node, upward_path = stack.pop()
        state = (node, tuple(upward_path))
        if state in seen_states:
            continue
        seen_states.add(state)

        for parent in parents_by_child.get(node, set()):
            if parent in upward_path:
                continue
            next_path = upward_path + [parent]
            if parent in managed_names:
                roots.add(parent)
                root_to_child = list(reversed(next_path))
                paths.append(" -> ".join(display_names.get(n, n) for n in root_to_child))
                continue
            stack.append((parent, next_path))

    root_list = sorted(display_names.get(name, name) for name in roots)
    unique_paths = list(dict.fromkeys(paths))
    return root_list, unique_paths


def _normalize_header(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.strip().lower())


def _resolve_header_indexes(header_row: list[Any]) -> dict[str, int] | None:
    header_map: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        normalized = _normalize_header(str(cell))
        if normalized:
            header_map[normalized] = idx

    aliases = {
        "section_no": ["sno", "sectionno", "slno", "serialno"],
        "section_title": ["section", "sectiontitle"],
        "module": ["module"],
        "sub_module": ["submodule", "submodules"],
        "feature_capability": ["featurecapability", "feature", "capability"],
        "description_details": ["descriptiondetails", "description", "details"],
    }

    resolved: dict[str, int] = {}
    for field_name, candidates in aliases.items():
        matched = next((header_map[c] for c in candidates if c in header_map), None)
        if matched is not None:
            resolved[field_name] = matched

    required = {"module", "sub_module", "feature_capability", "description_details"}
    if not required.issubset(set(resolved.keys())):
        return None
    return resolved


def _extract_excel_details(content: bytes) -> list[ReleaseDetailInput]:
    workbook = load_workbook(filename=BytesIO(content), data_only=True)
    sheet = workbook.active

    header_indexes: dict[str, int] | None = None
    details: list[ReleaseDetailInput] = []

    current_section_no: int | None = None
    current_section_title: str | None = None
    current_module: str | None = None
    current_sub_module: str | None = None

    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        if not any(cell is not None and str(cell).strip() for cell in values):
            continue

        if header_indexes is None:
            header_indexes = _resolve_header_indexes(values)
            if header_indexes is not None:
                normalized_header_values = {
                    _normalize_header(str(cell))
                    for cell in values
                    if cell is not None and str(cell).strip()
                }
                if normalized_header_values.intersection({"versionscope", "releaseversion", "scope"}):
                    raise HTTPException(
                        status_code=400,
                        detail="Scope/Version Scope column is no longer supported. Please remove it from the sheet.",
                    )
                continue

        if header_indexes is None:
            continue

        section_no_idx = header_indexes.get("section_no")
        section_title_idx = header_indexes.get("section_title")

        section_no_raw = values[section_no_idx] if section_no_idx is not None and section_no_idx < len(values) else None
        section_title_raw = values[section_title_idx] if section_title_idx is not None and section_title_idx < len(values) else None

        parsed_section_no: int | None = None
        if section_no_raw is not None and str(section_no_raw).strip():
            try:
                parsed_section_no = int(float(str(section_no_raw).strip()))
            except ValueError:
                parsed_section_no = None

        parsed_section_title = str(section_title_raw).strip() if section_title_raw is not None else ""
        if parsed_section_no is not None and parsed_section_title:
            current_section_no = parsed_section_no
            current_section_title = parsed_section_title
            current_module = None
            current_sub_module = None
            continue

        module_value = str(values[header_indexes["module"]] or "").strip() if header_indexes["module"] < len(values) else ""
        sub_module_value = str(values[header_indexes["sub_module"]] or "").strip() if header_indexes["sub_module"] < len(values) else ""
        feature_value = str(values[header_indexes["feature_capability"]] or "").strip() if header_indexes["feature_capability"] < len(values) else ""
        description_value = str(values[header_indexes["description_details"]] or "").strip() if header_indexes["description_details"] < len(values) else ""

        if _normalize_header(module_value) == "module" and _normalize_header(sub_module_value).startswith("submodule"):
            continue

        if module_value:
            current_module = module_value
        if sub_module_value:
            current_sub_module = sub_module_value

        if not feature_value and not description_value:
            continue

        if not feature_value:
            raise HTTPException(status_code=400, detail="Each detail row must include Feature / Capability.")

        details.append(
            ReleaseDetailInput(
                section_no=current_section_no,
                section_title=current_section_title,
                module=current_module,
                sub_module=current_sub_module,
                feature_capability=feature_value,
                description_details=description_value or None,
            )
        )

    if header_indexes is None:
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid sheet format. Required headers: Section No, Section, Module, Sub-Module, "
                "Feature/Capability, Description/Details"
            ),
        )

    if not details:
        raise HTTPException(status_code=400, detail="No release details found in uploaded sheet.")

    return details


def _extract_csv_details(content: bytes) -> list[ReleaseDetailInput]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    reader = csv.DictReader(StringIO(text))
    headers = reader.fieldnames or []
    normalized_headers = {_normalize_header(name): name for name in headers if name}

    required_aliases = {
        "module": ["module"],
        "sub_module": ["submodule", "submodules"],
        "feature_capability": ["featurecapability", "feature", "capability"],
        "description_details": ["descriptiondetails", "description", "details"],
    }
    optional_aliases = {
        "section_no": ["sno", "sectionno", "slno", "serialno"],
        "section_title": ["section", "sectiontitle"],
    }

    if set(normalized_headers.keys()).intersection({"versionscope", "releaseversion", "scope"}):
        raise HTTPException(
            status_code=400,
            detail="Scope/Version Scope column is no longer supported. Please remove it from the CSV.",
        )

    resolved_headers: dict[str, str] = {}
    for field_name, candidates in required_aliases.items():
        matched = next((normalized_headers[c] for c in candidates if c in normalized_headers), None)
        if matched is None:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Invalid CSV format. Required headers: Section No, Section, Module, Sub-Module, "
                    "Feature/Capability, Description/Details"
                ),
            )
        resolved_headers[field_name] = matched

    for field_name, candidates in optional_aliases.items():
        matched = next((normalized_headers[c] for c in candidates if c in normalized_headers), None)
        if matched is not None:
            resolved_headers[field_name] = matched

    details: list[ReleaseDetailInput] = []
    for row in reader:
        feature_value = str(row.get(resolved_headers["feature_capability"], "") or "").strip()
        if not feature_value:
            continue

        section_no: int | None = None
        section_no_header = resolved_headers.get("section_no")
        if section_no_header:
            raw_value = str(row.get(section_no_header, "") or "").strip()
            if raw_value:
                try:
                    section_no = int(float(raw_value))
                except ValueError:
                    section_no = None

        details.append(
            ReleaseDetailInput(
                section_no=section_no,
                section_title=(str(row.get(resolved_headers.get("section_title", ""), "") or "").strip() or None)
                if resolved_headers.get("section_title")
                else None,
                module=str(row.get(resolved_headers["module"], "") or "").strip() or None,
                sub_module=str(row.get(resolved_headers["sub_module"], "") or "").strip() or None,
                feature_capability=feature_value,
                description_details=str(row.get(resolved_headers["description_details"], "") or "").strip() or None,
            )
        )

    if not details:
        raise HTTPException(status_code=400, detail="No release details found in uploaded CSV.")

    return details


def _normalize_manual_details(details: list[ReleaseDetailInput]) -> list[ReleaseDetailInput]:
    if not details:
        raise HTTPException(status_code=400, detail="At least one manual detail row is required.")

    normalized: list[ReleaseDetailInput] = []
    for row in details:
        feature = row.feature_capability.strip()
        if not feature:
            raise HTTPException(status_code=400, detail="Feature / Capability is required for all manual rows.")

        normalized.append(
            ReleaseDetailInput(
                section_no=row.section_no,
                section_title=(row.section_title or "").strip() or None,
                module=(row.module or "").strip() or None,
                sub_module=(row.sub_module or "").strip() or None,
                feature_capability=feature,
                description_details=(row.description_details or "").strip() or None,
            )
        )

    return normalized


async def _create_release(
    *,
    session: DbSession,
    current_user: CurrentActiveUser,
    bump_type: BumpType,
    release_notes: str | None,
    details: list[ReleaseDetailInput] | None,
) -> dict[str, Any]:
    now = datetime.now(timezone.utc)
    today = now.date()

    active_release = (
        await session.exec(
            select(ProductRelease)
            .where(ProductRelease.end_date == ACTIVE_END_DATE)
            .order_by(ProductRelease.start_date.desc(), ProductRelease.created_at.desc())
        )
    ).first()

    if active_release:
        try:
            major, minor, patch = _parse_semver(active_release.version)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
    else:
        major, minor, patch = 0, 0, 0

    next_major, next_minor, next_patch = _bump(major, minor, patch, bump_type)
    next_version = f"{next_major}.{next_minor}.{next_patch}"

    existing = (await session.exec(select(ProductRelease).where(ProductRelease.version == next_version))).first()
    if existing is not None:
        raise HTTPException(status_code=409, detail=f"Release version '{next_version}' already exists.")

    if active_release is not None:
        active_release.end_date = today
        active_release.updated_at = now

    new_release = ProductRelease(
        version=next_version,
        major=next_major,
        minor=next_minor,
        patch=next_patch,
        release_notes=(release_notes or "").strip() or None,
        start_date=today,
        end_date=ACTIVE_END_DATE,
        created_by=current_user.id,
        created_at=now,
        updated_at=now,
    )
    session.add(new_release)
    await session.flush()

    if details:
        for idx, item in enumerate(details):
            session.add(
                ReleaseDetail(
                    release_id=new_release.id,
                    section_no=item.section_no,
                    section_title=item.section_title,
                    module=item.module,
                    sub_module=item.sub_module,
                    feature_capability=item.feature_capability,
                    description_details=item.description_details,
                    sort_order=idx,
                    created_at=now,
                )
            )

    try:
        current_packages = (
            await session.exec(
                select(Package)
                .where(
                    Package.end_date == ACTIVE_END_DATE,
                )
                .order_by(Package.name.asc(), Package.package_type.asc(), Package.synced_at.desc())
            )
        ).all()
    except ProgrammingError as exc:
        if "relation \"package\" does not exist" in str(exc):
            current_packages = []
        else:
            raise

    seen_keys: set[tuple[str, str, str]] = set()
    packages: list[Package] = []
    for pkg in current_packages:
        key = (pkg.service_name, pkg.name.lower(), pkg.package_type)
        if key in seen_keys:
            continue
        seen_keys.add(key)
        packages.append(pkg)

    packages_by_service: dict[str, list[Package]] = defaultdict(list)
    for pkg in packages:
        packages_by_service[pkg.service_name].append(pkg)

    filtered_packages: list[Package] = []
    for service_name, service_packages in packages_by_service.items():
        managed_rows = [pkg for pkg in service_packages if pkg.package_type == "managed"]
        transitive_rows = [pkg for pkg in service_packages if pkg.package_type == "transitive"]
        reachable_transitives = _compute_reachable_transitives_from_packages(
            managed_rows=managed_rows,
            transitive_rows=transitive_rows,
        )
        filtered_packages.extend(managed_rows)
        filtered_packages.extend(
            [pkg for pkg in transitive_rows if _normalize_pkg(pkg.name) in reachable_transitives]
        )

    for pkg in filtered_packages:
        session.add(
            ReleasePackageSnapshot(
                release_id=new_release.id,
                service_name=pkg.service_name,
                name=pkg.name,
                version=pkg.version,
                version_spec=pkg.version_spec,
                package_type=pkg.package_type,
                required_by=pkg.required_by,
                source=pkg.source,
                captured_at=now,
            )
        )
        pkg.release_id = new_release.id

    await session.commit()
    await session.refresh(new_release)

    return _release_to_payload(new_release, package_count=len(filtered_packages))


@router.get("")
@router.get("/")
async def get_releases(
    current_user: CurrentActiveUser,
    session: DbSession,
) -> list[dict[str, Any]]:
    releases = (
        await session.exec(
            select(ProductRelease).order_by(ProductRelease.start_date.desc(), ProductRelease.created_at.desc())
        )
    ).all()

    return [_release_to_payload(release) for release in releases]


@router.get("/current")
async def get_current_release(
    current_user: CurrentActiveUser,
    session: DbSession,
) -> dict[str, Any] | None:
    release = (
        await session.exec(
            select(ProductRelease)
            .where(ProductRelease.end_date == ACTIVE_END_DATE)
            .order_by(ProductRelease.start_date.desc(), ProductRelease.created_at.desc())
        )
    ).first()
    if release is None:
        return None
    return _release_to_payload(release)


@router.get("/{release_id}/details")
async def get_release_details(
    release_id: UUID,
    current_user: CurrentActiveUser,
    session: DbSession,
) -> list[dict[str, Any]]:
    release = await session.get(ProductRelease, release_id)
    if release is None:
        raise HTTPException(status_code=404, detail="Release not found")

    details = (
        await session.exec(
            select(ReleaseDetail)
            .where(ReleaseDetail.release_id == release_id)
            .order_by(ReleaseDetail.sort_order.asc(), ReleaseDetail.created_at.asc())
        )
    ).all()

    return [_detail_to_payload(item) for item in details]


@router.get("/{release_id}/packages")
async def get_release_packages(
    release_id: UUID,
    current_user: CurrentActiveUser,
    session: DbSession,
    service: str = Query(default="all"),
) -> list[dict[str, Any]]:
    release = await session.get(ProductRelease, release_id)
    if release is None:
        raise HTTPException(status_code=404, detail="Release not found")

    normalized_service = service.strip().lower()
    snapshots = (
        await session.exec(
            select(ReleasePackageSnapshot)
            .where(
                ReleasePackageSnapshot.release_id == release_id,
                *(()
                  if normalized_service == "all"
                  else (ReleasePackageSnapshot.service_name == normalized_service,)),
            )
            .order_by(
                ReleasePackageSnapshot.service_name.asc(),
                ReleasePackageSnapshot.package_type.asc(),
                ReleasePackageSnapshot.name.asc(),
            )
        )
    ).all()

    payloads: list[dict[str, Any]] = []
    snapshots_by_service: dict[str, list[ReleasePackageSnapshot]] = defaultdict(list)
    for row in snapshots:
        snapshots_by_service[row.service_name].append(row)

    for service_name, service_rows in snapshots_by_service.items():
        managed_rows = [row for row in service_rows if row.package_type == "managed"]
        transitive_rows = [row for row in service_rows if row.package_type == "transitive"]
        reachable_transitives = _compute_reachable_transitives_from_snapshot(
            managed_rows=managed_rows,
            transitive_rows=transitive_rows,
        )
        managed_names = {_normalize_pkg(row.name) for row in managed_rows}
        managed_by_norm = {_normalize_pkg(row.name): row for row in managed_rows}
        display_names = {_normalize_pkg(row.name): row.name for row in service_rows}
        parents_by_child = _build_release_parent_map(service_rows)

        for item in service_rows:
            payload = _release_package_to_payload(item)
            if item.package_type == "transitive":
                normalized_name = _normalize_pkg(item.name)
                if normalized_name not in reachable_transitives:
                    continue
                managed_roots, dependency_paths = _compute_managed_root_data_for_snapshot(
                    package_name=item.name,
                    parents_by_child=parents_by_child,
                    managed_names=managed_names,
                    display_names=display_names,
                )
                root_order = {root: idx for idx, root in enumerate(managed_roots)}
                dependency_paths = sorted(
                    dependency_paths,
                    key=lambda path: (
                        root_order.get(path.split(" -> ", 1)[0], 10_000),
                        path,
                    ),
                )
                managed_root_details = []
                for root_name in managed_roots:
                    root_row = managed_by_norm.get(_normalize_pkg(root_name))
                    if root_row is None:
                        continue
                    managed_root_details.append({"name": root_row.name, "version": root_row.version})
                payload["managed_roots"] = managed_roots
                payload["managed_root_details"] = managed_root_details
                payload["dependency_paths"] = dependency_paths
            else:
                payload["managed_roots"] = []
                payload["managed_root_details"] = []
                payload["dependency_paths"] = []
            payloads.append(payload)
    return payloads


@router.post("/bump")
async def bump_release(
    payload: ReleaseBumpRequest,
    current_user: CurrentActiveUser,
    session: DbSession,
) -> dict[str, Any]:
    await _require_release_permission(current_user, "publish_release")
    return await _create_release(
        session=session,
        current_user=current_user,
        bump_type=payload.bump_type,
        release_notes=payload.release_notes,
        details=None,
    )


@router.post("/bump-with-details")
async def bump_release_with_details(
    current_user: CurrentActiveUser,
    session: DbSession,
    bump_type: BumpType = Form(...),
    release_notes: str | None = Form(default=None),
    details_json: str | None = Form(default=None),
    details_file: UploadFile | None = File(default=None),
) -> dict[str, Any]:
    if details_file is None and (details_json is None or not details_json.strip()):
        raise HTTPException(status_code=400, detail="Provide either a sheet file or manual rows.")

    await _require_release_permission(current_user, "publish_release")

    parsed_details: list[ReleaseDetailInput]

    if details_file is not None:
        filename = (details_file.filename or "").lower()
        is_excel = filename.endswith(".xlsx") or filename.endswith(".xlsm") or filename.endswith(".xltx")
        is_csv = filename.endswith(".csv")

        if not (is_excel or is_csv):
            raise HTTPException(status_code=400, detail="Only .xlsx/.xlsm/.xltx/.csv files are supported.")

        content = await details_file.read()
        if not content:
            raise HTTPException(status_code=400, detail="Uploaded file is empty.")

        parsed_details = _extract_csv_details(content) if is_csv else _extract_excel_details(content)
    else:
        try:
            decoded = json.loads(details_json or "[]")
            manual_rows = [ReleaseDetailInput.model_validate(item) for item in decoded]
        except (json.JSONDecodeError, ValidationError, TypeError) as exc:
            raise HTTPException(status_code=400, detail="Invalid manual details payload.") from exc

        parsed_details = _normalize_manual_details(manual_rows)

    return await _create_release(
        session=session,
        current_user=current_user,
        bump_type=bump_type,
        release_notes=release_notes,
        details=parsed_details,
    )
