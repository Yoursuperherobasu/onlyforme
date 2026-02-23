from docx import Document

def create_rbac_tables_doc(filename="RBAC_Table_Design_Explanation.docx"):
    doc = Document()

    # Title
    doc.add_heading("RBAC Database Table Design Explanation", level=1)

    doc.add_paragraph(
        "This document explains the Role-Based Access Control (RBAC) database tables "
        "used in the platform. The focus is on how each table is designed, what data it "
        "stores, and how it is used during runtime authorization."
    )

    # ROLE TABLE
    doc.add_heading("1. role Table", level=2)
    doc.add_paragraph(
        "Purpose:\n"
        "The role table defines the list of application-level roles supported by the platform.\n\n"
        "What it stores:\n"
        "• Logical roles such as Developer, Tech Admin, Business Admin, and Tech Ops Admin.\n"
        "• These roles represent job functions within the application, not identity data.\n\n"
        "How it is used:\n"
        "• Roles are resolved dynamically at runtime based on Entra ID group mappings.\n"
        "• This table is static and changes rarely.\n\n"
        "Why it exists:\n"
        "• Avoids hardcoding roles in application logic.\n"
        "• Allows new roles to be added without code changes."
    )

    # PERMISSION TABLE
    doc.add_heading("2. permission Table", level=2)
    doc.add_paragraph(
        "Purpose:\n"
        "The permission table defines individual actions that can be performed in the platform.\n\n"
        "What it stores:\n"
        "• Fine-grained actions such as CREATE_FLOW, APPROVE_TECH, PUBLISH_FLOW, or REQUEST_MODEL.\n\n"
        "How it is used:\n"
        "• Permissions are checked at API level before executing sensitive operations.\n\n"
        "Why it exists:\n"
        "• Enables fine-grained access control.\n"
        "• Decouples business actions from user roles."
    )

    # ROLE_PERMISSION TABLE
    doc.add_heading("3. role_permission Table", level=2)
    doc.add_paragraph(
        "Purpose:\n"
        "The role_permission table defines which permissions are assigned to each role.\n\n"
        "What it stores:\n"
        "• Mapping between role IDs and permission IDs.\n\n"
        "How it is used:\n"
        "• When a role is resolved for a user, the platform fetches all permissions linked to that role.\n\n"
        "Data nature:\n"
        "• This is a configuration table populated manually during setup or via admin tools.\n\n"
        "Why it exists:\n"
        "• Allows permissions to be changed without redeploying code.\n"
        "• Supports scalable and maintainable RBAC."
    )

    # ENTRA GROUP ROLE MAP TABLE
    doc.add_heading("4. entra_group_role_map Table", level=2)
    doc.add_paragraph(
        "Purpose:\n"
        "This table maps Azure Entra ID groups to application roles.\n\n"
        "What it stores:\n"
        "• Entra ID group identifiers.\n"
        "• Corresponding application role IDs.\n\n"
        "How it is used:\n"
        "• At login, the platform reads group IDs from the Entra ID token.\n"
        "• These group IDs are matched against this table to determine application roles dynamically.\n\n"
        "Data nature:\n"
        "• Configuration-driven and rarely changes.\n\n"
        "Why it exists:\n"
        "• Keeps Entra ID as the identity source of truth.\n"
        "• Allows the application to control authorization independently.\n"
        "• Avoids storing roles directly on user records."
    )

    # Runtime Flow
    doc.add_heading("5. Runtime Authorization Flow", level=2)
    doc.add_paragraph(
        "1. User authenticates using Azure Entra ID.\n"
        "2. Authentication token includes Entra group IDs.\n"
        "3. Group IDs are mapped to roles using entra_group_role_map.\n"
        "4. Permissions are resolved using role_permission.\n"
        "5. API access is granted or denied based on permissions.\n\n"
        "No role or permission data is written to the database during login."
    )

    # Conclusion
    doc.add_heading("6. Summary", level=2)
    doc.add_paragraph(
        "This RBAC design ensures:\n"
        "• Dynamic authorization without storing user roles\n"
        "• Clear separation of identity and access control\n"
        "• Enterprise-ready governance and auditability\n"
        "• Flexibility to adapt to organizational changes"
    )

    doc.save(filename)
    print(f"RBAC table explanation document created: {filename}")

if __name__ == "__main__":
    create_rbac_tables_doc()
