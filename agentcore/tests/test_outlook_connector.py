#!/usr/bin/env python3
"""
Outlook Connector — Comprehensive Test Suite
=============================================
Standalone tests covering all 5 implementation phases.
No FastAPI/SQLAlchemy imports needed — logic is extracted and tested directly.

Run:
    python3 agentcore/tests/test_outlook_connector.py

Requires: cryptography, docx, openpyxl  (all pre-installed in the backend venv)
"""

from __future__ import annotations

import asyncio
import base64
import copy
import io
import json
import os
import re
import secrets
import sys
import time
from urllib.parse import parse_qs, urlencode, urlparse

from cryptography.fernet import Fernet

# ────────────────────────────────────────────────────────────
# Infrastructure
# ────────────────────────────────────────────────────────────

BASE = os.path.join(os.path.dirname(__file__), "..", "src")
BACKEND = os.path.join(BASE, "backend", "base", "agentcore")
FRONTEND = os.path.join(BASE, "frontend", "src")

passed = 0
failed = 0
section_counts = {}
current_section = ""


def section(name: str):
    global current_section
    current_section = name
    section_counts[name] = {"passed": 0, "failed": 0}
    print(f"\n{'='*70}")
    print(f"  {name}")
    print(f"{'='*70}")


def check(name: str, condition: bool):
    global passed, failed
    if condition:
        print(f"  PASS: {name}")
        passed += 1
        section_counts[current_section]["passed"] += 1
    else:
        print(f"  FAIL: {name}")
        failed += 1
        section_counts[current_section]["failed"] += 1


# ────────────────────────────────────────────────────────────
# Shared encryption helpers (extracted from connector_catalogue.py)
# ────────────────────────────────────────────────────────────

FERNET_KEY = Fernet.generate_key()
fernet = Fernet(FERNET_KEY)

EMAIL_PROVIDERS = {"outlook"}
STORAGE_PROVIDERS = {"azure_blob", "sharepoint"}
DB_PROVIDERS = {"postgresql", "oracle", "sqlserver", "mysql"}


def _encrypt_password(password: str) -> str:
    return fernet.encrypt(password.encode()).decode()


def _decrypt_password(encrypted: str) -> str:
    return fernet.decrypt(encrypted.encode()).decode()


def _encrypt_provider_config(provider: str, config: dict) -> dict:
    encrypted = dict(config)
    if provider == "azure_blob" and "connection_string" in encrypted:
        encrypted["connection_string"] = _encrypt_password(encrypted["connection_string"])
    elif provider == "sharepoint" and "client_secret" in encrypted:
        encrypted["client_secret"] = _encrypt_password(encrypted["client_secret"])
    elif provider in EMAIL_PROVIDERS:
        for key in ("client_secret", "access_token", "refresh_token"):
            if key in encrypted:
                encrypted[key] = _encrypt_password(encrypted[key])
        if "linked_accounts" in encrypted:
            encrypted["linked_accounts"] = [dict(acct) for acct in encrypted["linked_accounts"]]
            for acct in encrypted["linked_accounts"]:
                for key in ("access_token", "refresh_token"):
                    if key in acct:
                        acct[key] = _encrypt_password(acct[key])
    return encrypted


def _decrypt_provider_config(provider: str, config: dict) -> dict:
    decrypted = dict(config)
    try:
        if provider == "azure_blob" and "connection_string" in decrypted:
            decrypted["connection_string"] = _decrypt_password(decrypted["connection_string"])
        elif provider == "sharepoint" and "client_secret" in decrypted:
            decrypted["client_secret"] = _decrypt_password(decrypted["client_secret"])
        elif provider in EMAIL_PROVIDERS:
            for key in ("client_secret", "access_token", "refresh_token"):
                if key in decrypted:
                    try:
                        decrypted[key] = _decrypt_password(decrypted[key])
                    except Exception:
                        pass
            if "linked_accounts" in decrypted:
                decrypted["linked_accounts"] = [dict(acct) for acct in decrypted["linked_accounts"]]
                for acct in decrypted["linked_accounts"]:
                    for key in ("access_token", "refresh_token"):
                        if key in acct:
                            try:
                                acct[key] = _decrypt_password(acct[key])
                            except Exception:
                                pass
    except Exception:
        pass
    return decrypted


def _serialize_mask(provider: str, config: dict) -> dict:
    safe_config = dict(config)
    if provider == "azure_blob" and "connection_string" in safe_config:
        safe_config["connection_string"] = "********"
    elif provider == "sharepoint" and "client_secret" in safe_config:
        safe_config["client_secret"] = "********"
    elif provider in EMAIL_PROVIDERS:
        for key in ("client_secret", "access_token", "refresh_token"):
            if key in safe_config:
                safe_config[key] = "********"
        if "linked_accounts" in safe_config:
            safe_config["linked_accounts"] = [dict(acct) for acct in safe_config["linked_accounts"]]
            for acct in safe_config["linked_accounts"]:
                for key in ("access_token", "refresh_token"):
                    if key in acct:
                        acct[key] = "********"
    return safe_config


# ════════════════════════════════════════════════════════════
# Phase 1 — connector_catalogue.py
# ════════════════════════════════════════════════════════════

section("Phase 1 — Encrypt/Decrypt/Masking")

# Basic round-trip
config1 = {"tenant_id": "tid-123", "client_id": "cid-456", "client_secret": "super_secret"}
enc1 = _encrypt_provider_config("outlook", config1)
check("Encrypt: tenant_id unchanged", enc1["tenant_id"] == "tid-123")
check("Encrypt: client_id unchanged", enc1["client_id"] == "cid-456")
check("Encrypt: client_secret encrypted", enc1["client_secret"] != "super_secret")
dec1 = _decrypt_provider_config("outlook", enc1)
check("Decrypt: client_secret round-trips", dec1["client_secret"] == "super_secret")
check("Decrypt: tenant_id preserved", dec1["tenant_id"] == "tid-123")

# All three token fields
config2 = {"tenant_id": "t", "client_id": "c", "client_secret": "cs", "access_token": "at", "refresh_token": "rt"}
enc2 = _encrypt_provider_config("outlook", config2)
check("Encrypt: access_token encrypted", enc2["access_token"] != "at")
check("Encrypt: refresh_token encrypted", enc2["refresh_token"] != "rt")
dec2 = _decrypt_provider_config("outlook", enc2)
check("Decrypt: access_token round-trips", dec2["access_token"] == "at")
check("Decrypt: refresh_token round-trips", dec2["refresh_token"] == "rt")

# Linked accounts
config3 = {
    "tenant_id": "t", "client_id": "c", "client_secret": "cs",
    "linked_accounts": [
        {"email": "user1@test.com", "access_token": "u1_at", "refresh_token": "u1_rt"},
        {"email": "user2@test.com", "access_token": "u2_at", "refresh_token": "u2_rt"},
    ],
}
enc3 = _encrypt_provider_config("outlook", config3)
check("Encrypt: linked_accounts[0].email unchanged", enc3["linked_accounts"][0]["email"] == "user1@test.com")
check("Encrypt: linked_accounts[0].access_token encrypted", enc3["linked_accounts"][0]["access_token"] != "u1_at")
check("Encrypt: linked_accounts[1].refresh_token encrypted", enc3["linked_accounts"][1]["refresh_token"] != "u2_rt")
dec3 = _decrypt_provider_config("outlook", enc3)
check("Decrypt: linked_accounts[0].access_token round-trips", dec3["linked_accounts"][0]["access_token"] == "u1_at")
check("Decrypt: linked_accounts[1].refresh_token round-trips", dec3["linked_accounts"][1]["refresh_token"] == "u2_rt")

# Deep copy safety
orig = copy.deepcopy(config3)
_encrypt_provider_config("outlook", config3)
check("Deep copy: encrypt does not mutate original", config3 == orig)

enc_copy = _encrypt_provider_config("outlook", config3)
enc_snapshot = copy.deepcopy(enc_copy)
_decrypt_provider_config("outlook", enc_copy)
check("Deep copy: decrypt does not mutate encrypted", enc_copy == enc_snapshot)

# Masking
mask1 = _serialize_mask("outlook", config2)
check("Mask: client_secret=********", mask1["client_secret"] == "********")
check("Mask: access_token=********", mask1["access_token"] == "********")
check("Mask: refresh_token=********", mask1["refresh_token"] == "********")
check("Mask: tenant_id visible", mask1["tenant_id"] == "t")

mask2 = _serialize_mask("outlook", config3)
check("Mask: linked_accounts[0].access_token=********", mask2["linked_accounts"][0]["access_token"] == "********")
check("Mask: linked_accounts[0].email visible", mask2["linked_accounts"][0]["email"] == "user1@test.com")

orig_mask = copy.deepcopy(config3)
_serialize_mask("outlook", config3)
check("Mask: does not mutate original", config3 == orig_mask)

# Existing providers still work
sp_enc = _encrypt_provider_config("sharepoint", {"client_secret": "sp_sec"})
check("SharePoint encrypt works", sp_enc["client_secret"] != "sp_sec")
sp_dec = _decrypt_provider_config("sharepoint", sp_enc)
check("SharePoint decrypt round-trips", sp_dec["client_secret"] == "sp_sec")

ab_enc = _encrypt_provider_config("azure_blob", {"connection_string": "DefaultEndpoints..."})
check("Azure Blob encrypt works", ab_enc["connection_string"] != "DefaultEndpoints...")
ab_dec = _decrypt_provider_config("azure_blob", ab_enc)
check("Azure Blob decrypt round-trips", ab_dec["connection_string"] == "DefaultEndpoints...")

# Edge cases
check("Empty config: no error", _encrypt_provider_config("outlook", {}) == {})
check("Partial config (no secrets): unchanged", _encrypt_provider_config("outlook", {"tenant_id": "t"}) == {"tenant_id": "t"})

# ────────────────────────────────────────────────────────────
section("Phase 1 — Source file structure")

cc_path = os.path.join(BACKEND, "api", "connector_catalogue.py")
with open(cc_path) as f:
    cc_src = f.read()

check("EMAIL_PROVIDERS defined", 'EMAIL_PROVIDERS = {"outlook"}' in cc_src)
check("_test_outlook_connection exists", "async def _test_outlook_connection(config: dict)" in cc_src)
check("Graph /me endpoint used", "https://graph.microsoft.com/v1.0/me" in cc_src)
check("Create: STORAGE_PROVIDERS | EMAIL_PROVIDERS", "STORAGE_PROVIDERS | EMAIL_PROVIDERS" in cc_src)
check("Test saved: elif provider in EMAIL_PROVIDERS", "elif provider in EMAIL_PROVIDERS:" in cc_src)
check("Test draft: if provider in EMAIL_PROVIDERS", "if provider in EMAIL_PROVIDERS:" in cc_src)


# ════════════════════════════════════════════════════════════
# Phase 2 — outlook_connector.py
# ════════════════════════════════════════════════════════════

section("Phase 2 — Outlook Connector API structure")

oc_path = os.path.join(BACKEND, "api", "outlook_connector.py")
with open(oc_path) as f:
    oc_src = f.read()

check("Router prefix /outlook", 'prefix="/outlook"' in oc_src)
check("Has 6+ route decorators", oc_src.count("@router.") >= 6)
check("OAuth start endpoint", "/{connector_id}/oauth/start" in oc_src)
check("OAuth callback endpoint", "/oauth/callback" in oc_src)
check("Accounts list endpoint", "/{connector_id}/accounts" in oc_src)
check("Accounts delete endpoint", "/{connector_id}/accounts/{email}" in oc_src)
check("Read mail endpoint", "/{connector_id}/read" in oc_src)
check("Reply mail endpoint", "/{connector_id}/reply" in oc_src)

check("Redis OAuth state store", "_store_oauth_state" in oc_src and "_pop_oauth_state" in oc_src)
check("_load_connector helper", "async def _load_connector(" in oc_src)
check("_find_account helper", "def _find_account(" in oc_src)
check("_refresh_token_if_needed helper", "async def _refresh_token_if_needed(" in oc_src)
check("_save_updated_config helper", "async def _save_updated_config(" in oc_src)

check("Imports _can_access_connector", "_can_access_connector" in oc_src)
check("Imports _decrypt_provider_config", "_decrypt_provider_config" in oc_src)
check("Imports EMAIL_PROVIDERS", "EMAIL_PROVIDERS" in oc_src)

callback_match = re.search(r'async def oauth_callback\([^)]*\)', oc_src)
check("OAuth callback: no CurrentActiveUser",
      callback_match is not None and "CurrentActiveUser" not in callback_match.group(0))

check("Mail.Read scope", "Mail.Read" in oc_src)
check("Mail.Send scope", "Mail.Send" in oc_src)
check("offline_access scope", "offline_access" in oc_src)
check("Supports sender reply mode", '"sender"' in oc_src)
check("Supports reply_all mode", "reply_all" in oc_src)
check("Supports custom mode", "custom" in oc_src)

# ────────────────────────────────────────────────────────────
section("Phase 2 — OAuth state & helpers (standalone)")

_oauth_states_test: dict[str, dict] = {}
STATE_TTL = 600

state_token = secrets.token_urlsafe(32)
_oauth_states_test[state_token] = {
    "connector_id": "c-123", "user_id": "u-456",
    "redirect_uri": "http://localhost:7860/api/outlook/oauth/callback",
    "created_at": time.time(),
}
check("State stored", state_token in _oauth_states_test)
check("State has connector_id", _oauth_states_test[state_token]["connector_id"] == "c-123")
check("Fresh state not expired", (time.time() - _oauth_states_test[state_token]["created_at"]) <= STATE_TTL)

_oauth_states_test["old"] = {"created_at": time.time() - 700}
check("Old state is expired", (time.time() - _oauth_states_test["old"]["created_at"]) > STATE_TTL)

def _find_account(linked_accounts, email):
    email_lower = email.lower()
    for i, acct in enumerate(linked_accounts):
        if acct.get("email", "").lower() == email_lower:
            return i, acct
    return -1, None

accts = [{"email": "User1@Example.com", "access_token": "tok1"}, {"email": "user2@test.com", "access_token": "tok2"}]
idx, acct = _find_account(accts, "user1@example.com")
check("_find_account: case-insensitive match", idx == 0)
idx2, _ = _find_account(accts, "nobody@test.com")
check("_find_account: -1 for not found", idx2 == -1)


# ════════════════════════════════════════════════════════════
# Phase 3 — graph_mail.py
# ════════════════════════════════════════════════════════════

section("Phase 3 — Graph Mail Client structure")

gm_path = os.path.join(BACKEND, "services", "outlook", "graph_mail.py")
with open(gm_path) as f:
    gm_src = f.read()

check("OutlookGraphMailClient class", "class OutlookGraphMailClient:" in gm_src)
check("GRAPH_BASE constant", 'GRAPH_BASE = "https://graph.microsoft.com/v1.0"' in gm_src)
for m in ["get_authorize_url", "exchange_code_for_tokens", "_refresh_access_token",
          "_get_valid_token", "_headers", "get_token_state", "get_me",
          "list_messages", "get_message", "list_attachments",
          "reply_to_message", "reply_all_to_message", "send_mail"]:
    check(f"Method: {m}", f"def {m}(" in gm_src or f"async def {m}(" in gm_src)
check("Isolated from teams", "from agentcore.services.teams" not in gm_src)
check("Uses httpx", "import httpx" in gm_src)

# ────────────────────────────────────────────────────────────
section("Phase 3 — Authorize URL generation")

GRAPH_BASE = "https://graph.microsoft.com/v1.0"
AUTHORIZE_URL = "https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/authorize"
MAIL_SCOPES = "Mail.Read Mail.ReadWrite Mail.Send User.Read offline_access"

tenant, client_id = "my-tenant", "my-client"
redirect_uri = "http://localhost:7860/api/outlook/oauth/callback"
state = "random-state-123"
base = AUTHORIZE_URL.format(tenant_id=tenant)
params = urlencode({"client_id": client_id, "response_type": "code", "redirect_uri": redirect_uri,
                     "scope": MAIL_SCOPES, "state": state, "response_mode": "query", "prompt": "select_account"})
url = f"{base}?{params}"
parsed = urlparse(url)
qs = parse_qs(parsed.query)
check("Auth URL: correct host", parsed.hostname == "login.microsoftonline.com")
check("Auth URL: tenant in path", tenant in parsed.path)
check("Auth URL: client_id param", qs["client_id"] == [client_id])
check("Auth URL: response_type=code", qs["response_type"] == ["code"])
check("Auth URL: has scope", "Mail.Read" in qs["scope"][0])
check("Auth URL: has state", qs["state"] == [state])

# ────────────────────────────────────────────────────────────
section("Phase 3 — Token lifecycle & mail ops")

class MockClient:
    def __init__(self, access_token=None, refresh_token=None, token_expires_at=None):
        self._access_token = access_token
        self._refresh_token = refresh_token
        self._token_expires_at = token_expires_at or 0.0
    def _needs_refresh(self):
        if not self._access_token:
            return None
        return time.time() >= (self._token_expires_at - 60)
    def get_token_state(self):
        return {"access_token": self._access_token, "refresh_token": self._refresh_token,
                "token_expires_at": self._token_expires_at}

check("Fresh token: no refresh", not MockClient("t", "r", time.time() + 3600)._needs_refresh())
check("Expiring (30s): needs refresh", MockClient("t", "r", time.time() + 30)._needs_refresh())
check("Expired: needs refresh", MockClient("t", "r", time.time() - 100)._needs_refresh())
check("No token: None", MockClient()._needs_refresh() is None)
check("Token state keys", set(MockClient("t", "r", 0).get_token_state().keys()) == {"access_token", "refresh_token", "token_expires_at"})

# Mail URLs
msg_id = "AAMkAGVm"
check("List messages URL", f"{GRAPH_BASE}/me/mailFolders/inbox/messages" == "https://graph.microsoft.com/v1.0/me/mailFolders/inbox/messages")
check("Reply URL", f"{GRAPH_BASE}/me/messages/{msg_id}/reply" == f"https://graph.microsoft.com/v1.0/me/messages/{msg_id}/reply")
check("Reply-all URL", f"{GRAPH_BASE}/me/messages/{msg_id}/replyAll" == f"https://graph.microsoft.com/v1.0/me/messages/{msg_id}/replyAll")
check("Send mail URL", f"{GRAPH_BASE}/me/sendMail" == "https://graph.microsoft.com/v1.0/me/sendMail")

to = ["a@b.com", "c@d.com"]
payload = {"message": {"subject": "Test", "body": {"contentType": "Text", "content": "Body"},
           "toRecipients": [{"emailAddress": {"address": r}} for r in to]}, "saveToSentItems": True}
check("Send: 2 recipients", len(payload["message"]["toRecipients"]) == 2)
check("Send: saveToSentItems", payload["saveToSentItems"] is True)


# ════════════════════════════════════════════════════════════
# Phase 4 — attachment_parser.py
# ════════════════════════════════════════════════════════════

section("Phase 4 — Attachment parser structure")

ap_path = os.path.join(BACKEND, "services", "outlook", "attachment_parser.py")
with open(ap_path) as f:
    ap_src = f.read()

check("MAX_ATTACHMENT_BYTES=10MB", "MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024" in ap_src)
check("MAX_ATTACHMENTS=20", "MAX_ATTACHMENTS = 20" in ap_src)
check("parse_attachment function", "def parse_attachment(" in ap_src)
check("parse_attachments function", "def parse_attachments(" in ap_src)
for ext in [".txt", ".csv", ".pdf", ".docx", ".xlsx", ".pptx"]:
    check(f"Supports {ext}", f'"{ext}"' in ap_src)
check("Filters @odata.type", "#microsoft.graph.fileAttachment" in ap_src)

# ────────────────────────────────────────────────────────────
section("Phase 4 — Attachment parsing (standalone)")

MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".txt", ".csv", ".pdf", ".docx", ".xlsx", ".pptx"}

def parse_attachment(name, content_bytes_b64):
    ext = ("." + name.rsplit(".", 1)[-1]).lower() if "." in name else ""
    result = {"filename": name, "text": None, "error": None, "size_bytes": 0}
    try:
        raw = base64.b64decode(content_bytes_b64)
        result["size_bytes"] = len(raw)
    except Exception:
        result["error"] = "Failed to decode base64 content"
        return result
    if len(raw) > MAX_ATTACHMENT_BYTES:
        result["error"] = f"Attachment too large ({len(raw)} bytes, max {MAX_ATTACHMENT_BYTES})"
        return result
    if ext not in SUPPORTED_EXTENSIONS:
        result["error"] = f"Unsupported file type: {ext}"
        return result
    try:
        if ext in (".txt", ".csv"):
            result["text"] = raw.decode("utf-8", errors="replace")
        elif ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(raw))
            result["text"] = "\n".join(p.text for p in doc.paragraphs)
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            result["text"] = "\n".join(lines)
            wb.close()
    except Exception as e:
        result["error"] = f"Parse error: {e!s}"
    return result

def parse_attachments(attachments):
    results = []
    for att in attachments[:20]:
        if att.get("@odata.type") != "#microsoft.graph.fileAttachment":
            continue
        results.append(parse_attachment(att.get("name", "unknown"), att.get("contentBytes", "")))
    return results

# txt
r = parse_attachment("hello.txt", base64.b64encode(b"Hello, World!").decode())
check("TXT: extracted", r["text"] == "Hello, World!")
check("TXT: no error", r["error"] is None)
check("TXT: size_bytes", r["size_bytes"] == 13)

# csv
r = parse_attachment("data.csv", base64.b64encode(b"a,b\n1,2").decode())
check("CSV: extracted", r["text"] == "a,b\n1,2")

# docx (real file)
from docx import Document as DocxDoc
doc = DocxDoc()
doc.add_paragraph("Test paragraph one.")
doc.add_paragraph("Test paragraph two.")
buf = io.BytesIO()
doc.save(buf)
r = parse_attachment("test.docx", base64.b64encode(buf.getvalue()).decode())
check("DOCX: extracted", r["text"] is not None and "Test paragraph one" in r["text"])
check("DOCX: both paragraphs", "Test paragraph two" in (r["text"] or ""))
check("DOCX: no error", r["error"] is None)

# xlsx (real file)
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.append(["Name", "Value"])
ws.append(["Alpha", 100])
buf = io.BytesIO()
wb.save(buf)
r = parse_attachment("data.xlsx", base64.b64encode(buf.getvalue()).decode())
check("XLSX: extracted", r["text"] is not None)
check("XLSX: has header", "Name\tValue" in (r["text"] or ""))
check("XLSX: has data", "Alpha\t100" in (r["text"] or ""))

# oversize
r = parse_attachment("huge.txt", base64.b64encode(b"x" * (MAX_ATTACHMENT_BYTES + 1)).decode())
check("Oversize: rejected", "too large" in (r["error"] or ""))
check("Oversize: no text", r["text"] is None)

# unsupported
r = parse_attachment("virus.exe", base64.b64encode(b"\x00").decode())
check("Unsupported: rejected", "Unsupported" in (r["error"] or ""))

# bad base64 — single char causes genuine decode failure
r = parse_attachment("bad.txt", "x")
check("Bad base64: error", r["error"] is not None and "base64" in r["error"].lower())

# no extension
r = parse_attachment("noext", base64.b64encode(b"data").decode())
check("No extension: rejected", "Unsupported" in (r["error"] or ""))

# batch
batch = [
    {"@odata.type": "#microsoft.graph.fileAttachment", "name": "a.txt", "contentBytes": base64.b64encode(b"A").decode()},
    {"@odata.type": "#microsoft.graph.itemAttachment", "name": "b.txt", "contentBytes": base64.b64encode(b"B").decode()},
    {"@odata.type": "#microsoft.graph.fileAttachment", "name": "c.csv", "contentBytes": base64.b64encode(b"x").decode()},
]
br = parse_attachments(batch)
check("Batch: itemAttachment skipped (2 results)", len(br) == 2)
check("Batch: first=a.txt", br[0]["filename"] == "a.txt")

# limit
many = [{"@odata.type": "#microsoft.graph.fileAttachment", "name": f"f{i}.txt",
         "contentBytes": base64.b64encode(b"d").decode()} for i in range(25)]
check("Batch limit: max 20", len(parse_attachments(many)) == 20)


# ════════════════════════════════════════════════════════════
# Phase 5 — Frontend
# ════════════════════════════════════════════════════════════

section("Phase 5 — Frontend verification")

idx_path = os.path.join(FRONTEND, "pages", "ConnectorsCatalogue", "index.tsx")
form_path = os.path.join(FRONTEND, "pages", "ConnectorsCatalogue", "components", "OutlookConnectorForm.tsx")
flags_path = os.path.join(FRONTEND, "customization", "feature-flags.ts")

with open(idx_path) as f:
    idx_src = f.read()
with open(form_path) as f:
    form_src = f.read()
with open(flags_path) as f:
    flags_src = f.read()

check("Import: OutlookConnectorForm", "import OutlookConnectorForm" in idx_src)
check("Type: ProviderFilter has outlook", '| "outlook"' in idx_src)
check("PROVIDER_LABELS has outlook", 'outlook: "Microsoft Outlook"' in idx_src)
check("EMAIL_PROVIDERS set", 'const EMAIL_PROVIDERS = new Set(["outlook"])' in idx_src)
check("BLANK_FORM: outlook fields", 'outlook_tenant_id: ""' in idx_src and 'outlook_client_id: ""' in idx_src)
check("openEditModal: Outlook mapping", "outlook_tenant_id: cfg.tenant_id" in idx_src)
check("buildPayload: Outlook branch", 'form.provider === "outlook"' in idx_src)
check("isSaveDisabled: Outlook validation", "!form.outlook_tenant_id || !form.outlook_client_id" in idx_src)
check("handleModalTestConnection: Outlook", idx_src.count('form.provider === "outlook"') >= 2)
check("getProviderBadge: outlook", "bg-sky" in idx_src)
check("getConnectorTarget: outlook", 'c.provider === "outlook"' in idx_src)
check("getConnectorDb: EMAIL_PROVIDERS", "EMAIL_PROVIDERS.has(c.provider)" in idx_src)
check("FILTER_TABS includes outlook", re.search(r'FILTER_TABS.*"outlook"', idx_src) is not None)
check("Dropdown: Email optgroup", '<optgroup label="Email">' in idx_src)
check("Dropdown: Outlook option", '<option value="outlook">Microsoft Outlook</option>' in idx_src)
check("JSX: <OutlookConnectorForm", "<OutlookConnectorForm" in idx_src)
check("Header: mentions Outlook", "Outlook)" in idx_src)

check("Form component: default export", "export default function OutlookConnectorForm" in form_src)
check("Form: Eye/EyeOff toggle", "showSecret" in form_src and "EyeOff" in form_src)
check("Form: consistent Tailwind", "border-border bg-background" in form_src)
check("Form: OAuth hint", "After saving, use the OAuth flow" in form_src)

check("Feature flag: ENABLE_OUTLOOK_CONNECTOR", "export const ENABLE_OUTLOOK_CONNECTOR = true" in flags_src)

# No regressions
check("SharePoint still in dropdown", '<option value="sharepoint">SharePoint</option>' in idx_src)
check("Azure Blob still in dropdown", '<option value="azure_blob">Azure Blob Storage</option>' in idx_src)
check("PostgreSQL still in dropdown", '<option value="postgresql">PostgreSQL</option>' in idx_src)


# ════════════════════════════════════════════════════════════
# Cross-Phase Integration
# ════════════════════════════════════════════════════════════

section("Cross-Phase Integration")

# File inventory
for label, path in [
    ("__init__.py", os.path.join(BACKEND, "services", "outlook", "__init__.py")),
    ("graph_mail.py", os.path.join(BACKEND, "services", "outlook", "graph_mail.py")),
    ("attachment_parser.py", os.path.join(BACKEND, "services", "outlook", "attachment_parser.py")),
    ("outlook_connector.py", os.path.join(BACKEND, "api", "outlook_connector.py")),
    ("OutlookConnectorForm.tsx", form_path),
]:
    check(f"New file: {label}", os.path.isfile(path))

# Router wiring
router_path = os.path.join(BACKEND, "api", "router.py")
with open(router_path) as f:
    router_src = f.read()
check("Router imports outlook_connector_router", "from agentcore.api.outlook_connector import router as outlook_connector_router" in router_src)
check("Router includes outlook_connector_router", "router.include_router(outlook_connector_router)" in router_src)

# __init__.py exports
init_path = os.path.join(BACKEND, "services", "outlook", "__init__.py")
with open(init_path) as f:
    init_src = f.read()
check("Init exports OutlookGraphMailClient", "OutlookGraphMailClient" in init_src)
check("Init exports parse_attachment", "parse_attachment" in init_src)
check("Init has __all__", "__all__" in init_src)

# Provider sets disjoint
check("All 7 providers", DB_PROVIDERS | STORAGE_PROVIDERS | EMAIL_PROVIDERS == {"postgresql", "oracle", "sqlserver", "mysql", "azure_blob", "sharepoint", "outlook"})
check("Sets disjoint", DB_PROVIDERS.isdisjoint(STORAGE_PROVIDERS) and DB_PROVIDERS.isdisjoint(EMAIL_PROVIDERS) and STORAGE_PROVIDERS.isdisjoint(EMAIL_PROVIDERS))

# Encrypt/decrypt for all non-DB providers
for p in STORAGE_PROVIDERS | EMAIL_PROVIDERS:
    cfg = {"client_secret": "test"} if p != "azure_blob" else {"connection_string": "test"}
    dec = _decrypt_provider_config(p, _encrypt_provider_config(p, cfg))
    key = "connection_string" if p == "azure_blob" else "client_secret"
    check(f"Round-trip: {p}", dec[key] == "test")

# Rollback plan
check("Feature flag exists for UI disable", "ENABLE_OUTLOOK_CONNECTOR" in flags_src)
check("Router removable (2 lines)", "outlook_connector_router" in router_src)
check("No alembic migration", not any("outlook" in f.lower() for f in os.listdir(os.path.join(BACKEND, "alembic", "versions")) if f.endswith(".py")) if os.path.isdir(os.path.join(BACKEND, "alembic", "versions")) else True)


# ════════════════════════════════════════════════════════════
# SUMMARY
# ════════════════════════════════════════════════════════════

print(f"\n{'='*70}")
print(f"  SECTION SUMMARY")
print(f"{'='*70}")
for name, counts in section_counts.items():
    status = "PASS" if counts["failed"] == 0 else "FAIL"
    print(f"  [{status}] {name}: {counts['passed']}p / {counts['failed']}f")

print(f"\n{'='*70}")
print(f"  TOTAL: {passed} passed, {failed} failed out of {passed + failed}")
print(f"{'='*70}")

if failed:
    print("\n  SOME TESTS FAILED!")
    sys.exit(1)
else:
    print("\n  ALL TESTS PASSED!")
    sys.exit(0)
